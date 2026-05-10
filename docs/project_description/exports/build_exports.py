from pathlib import Path
import re, html
import markdown
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib import colors

root = Path('/opt/hermes/Destabilization_Signal_Monitor/docs/project_description')
export = root / 'exports'
export.mkdir(parents=True, exist_ok=True)

docs = [
    'scientific_paper.md',
    'scientific_paper_formal.md',
    'executive_summary_de.md',
    'detailed_documentation.md',
]


def build_html(md_path: Path):
    text = md_path.read_text(encoding='utf-8')
    body = markdown.markdown(text, extensions=['tables', 'fenced_code'])
    title = md_path.stem
    html_text = f'''<!DOCTYPE html>
<html lang="de">
<head>
<meta charset="utf-8" />
<title>{html.escape(title)}</title>
<style>
body {{ font-family: Arial, Helvetica, sans-serif; line-height: 1.5; color: #0f172a; max-width: 980px; margin: 40px auto; padding: 0 24px; }}
h1,h2,h3,h4 {{ color: #0f172a; }}
h1 {{ border-bottom: 2px solid #cbd5e1; padding-bottom: 8px; }}
h2 {{ margin-top: 32px; border-bottom: 1px solid #e2e8f0; padding-bottom: 4px; }}
img {{ max-width: 100%; border: 1px solid #cbd5e1; border-radius: 8px; background: #fff; padding: 8px; }}
table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
th, td {{ border: 1px solid #cbd5e1; padding: 8px 10px; vertical-align: top; }}
th {{ background: #f8fafc; text-align: left; }}
code {{ background: #f1f5f9; padding: 2px 4px; border-radius: 4px; }}
pre code {{ display: block; padding: 12px; overflow-x: auto; }}
blockquote {{ border-left: 4px solid #94a3b8; margin-left: 0; padding-left: 12px; color: #334155; }}
.smallnote {{ margin-top: 36px; font-size: 12px; color: #475569; }}
@media print {{ body {{ max-width: none; margin: 0; }} img {{ page-break-inside: avoid; }} h2, h3 {{ page-break-after: avoid; }} }}
</style>
</head>
<body>
{body}
<div class="smallnote">Exportfaehige HTML-Fassung, automatisch aus Markdown erzeugt.</div>
</body>
</html>'''
    (export / f'{md_path.stem}.html').write_text(html_text, encoding='utf-8')


def md_to_blocks(text: str):
    lines = text.splitlines()
    blocks = []
    current = []
    for line in lines:
        if line.strip().startswith('!['):
            if current:
                blocks.append(('paragraph', '\n'.join(current).strip()))
                current = []
            blocks.append(('image_ref', line.strip()))
            continue
        if not line.strip():
            if current:
                blocks.append(('paragraph', '\n'.join(current).strip()))
                current = []
            continue
        if line.startswith('#'):
            if current:
                blocks.append(('paragraph', '\n'.join(current).strip()))
                current = []
            level = len(line) - len(line.lstrip('#'))
            blocks.append((f'h{level}', line[level:].strip()))
        else:
            current.append(line)
    if current:
        blocks.append(('paragraph', '\n'.join(current).strip()))
    return blocks


def clean_inline(s: str) -> str:
    s = re.sub(r'`([^`]+)`', r'\1', s)
    s = re.sub(r'!\[([^\]]*)\]\(([^)]+)\)', r'[Abbildung: \1 -> \2]', s)
    s = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', s)
    s = s.replace('\\[', '').replace('\\]', '')
    return s


def build_docx(md_path: Path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10.5)
    blocks = md_to_blocks(md_path.read_text(encoding='utf-8'))
    for kind, content in blocks:
        if kind == 'h1':
            doc.add_heading(clean_inline(content), level=0)
        elif kind == 'h2':
            doc.add_heading(clean_inline(content), level=1)
        elif kind == 'h3':
            doc.add_heading(clean_inline(content), level=2)
        elif kind == 'h4':
            doc.add_heading(clean_inline(content), level=3)
        elif kind == 'image_ref':
            p = doc.add_paragraph()
            p.add_run(clean_inline(content)).italic = True
        else:
            txt = clean_inline(content)
            p = doc.add_paragraph(txt)
            p.paragraph_format.space_after = Pt(6)
    doc.add_paragraph('Automatisch erzeugte DOCX-Exportfassung. SVG-Abbildungen bleiben ueber Referenzen auf die Diagrammdateien nachvollziehbar.')
    doc.save(export / f'{md_path.stem}.docx')


def build_pdf(md_path: Path):
    pdf_path = export / f'{md_path.stem}.pdf'
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=1.7*cm, bottomMargin=1.7*cm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='Body2', parent=styles['BodyText'], fontName='Helvetica', fontSize=10, leading=14, spaceAfter=8))
    styles.add(ParagraphStyle(name='H1x', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=colors.HexColor('#0f172a'), spaceBefore=10, spaceAfter=10))
    styles.add(ParagraphStyle(name='H2x', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=14, leading=18, textColor=colors.HexColor('#0f172a'), spaceBefore=8, spaceAfter=8))
    styles.add(ParagraphStyle(name='H3x', parent=styles['Heading3'], fontName='Helvetica-Bold', fontSize=12, leading=16, textColor=colors.HexColor('#0f172a'), spaceBefore=6, spaceAfter=6))
    story = []
    for kind, content in md_to_blocks(md_path.read_text(encoding='utf-8')):
        txt = html.escape(clean_inline(content)).replace('\n', '<br/>')
        if kind == 'h1':
            story.append(Paragraph(txt, styles['H1x']))
        elif kind == 'h2':
            story.append(Paragraph(txt, styles['H2x']))
        elif kind == 'h3' or kind == 'h4':
            story.append(Paragraph(txt, styles['H3x']))
        elif kind == 'image_ref':
            story.append(Paragraph(f'<i>{txt}</i>', styles['Body2']))
        else:
            story.append(Paragraph(txt, styles['Body2']))
        story.append(Spacer(1, 0.1*cm))
    story.append(Spacer(1, 0.3*cm))
    story.append(Paragraph('Automatisch erzeugte PDF-Exportfassung. SVG-Abbildungen werden in dieser Minimalfassung als referenzierte Diagrammhinweise gefuehrt.', styles['Body2']))
    doc.build(story)


for name in docs:
    p = root / name
    build_html(p)
    build_docx(p)
    build_pdf(p)
print('generated html/docx/pdf exports for ' + ', '.join(docs))
